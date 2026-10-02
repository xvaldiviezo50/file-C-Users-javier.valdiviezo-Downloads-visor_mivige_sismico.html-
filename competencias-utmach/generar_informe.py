from docx import Document
from docx.shared import Inches,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import matplotlib.pyplot as plt, numpy as np, os
os.makedirs("competencias-utmach/report_assets",exist_ok=True)
names=["Tecnológica","Pedagógica","Comunicativa","Gestión","Investigativa"]
E=np.array([15.44,19.85,24.26,31.62,24.26]); I=np.array([11.76,5.15,17.65,9.56,7.35]); N=np.array([72.79,75,58.09,58.82,68.38]); A=np.array([4.30,4.25,4.05,3.93,4.21])
# gráficos
plt.figure(figsize=(9,5)); y=np.arange(5); plt.barh(y,E,label="Explorador");plt.barh(y,I,left=E,label="Integrador");plt.barh(y,N,left=E+I,label="Innovador");plt.yticks(y,names);plt.xlim(0,100);plt.xlabel("% de docentes");plt.title("Distribución por competencia");plt.legend(ncol=3,loc="lower center",bbox_to_anchor=(.5,-.25));plt.tight_layout();plt.savefig("competencias-utmach/report_assets/distribucion.png",dpi=180,bbox_inches="tight");plt.close()
ang=np.linspace(0,2*np.pi,5,endpoint=False).tolist(); vals=N.tolist()+[N[0]]; aa=ang+[ang[0]]; fig=plt.figure(figsize=(6.5,6));ax=fig.add_subplot(111,polar=True);ax.plot(aa,vals,lw=2);ax.fill(aa,vals,alpha=.15);ax.set_xticks(ang);ax.set_xticklabels(names);ax.set_ylim(0,100);ax.set_title("Perfil institucional: nivel Innovador",pad=20);plt.tight_layout();plt.savefig("competencias-utmach/report_assets/pentagono.png",dpi=180,bbox_inches="tight");plt.close()
plt.figure(figsize=(8,4.3));b=plt.bar(names,A);plt.ylim(0,5);plt.ylabel("Promedio /5");plt.title("Promedio global por competencia");plt.xticks(rotation=18,ha="right");[plt.text(x.get_x()+x.get_width()/2,v+.05,f"{v:.2f}",ha="center") for x,v in zip(b,A)];plt.tight_layout();plt.savefig("competencias-utmach/report_assets/promedios.png",dpi=180,bbox_inches="tight");plt.close()
ages=["24–35","36–45","46–55","56–65",">66"]; an=[25,47,42,19,3];plt.figure(figsize=(7.5,4));b=plt.bar(ages,an);plt.ylabel("Docentes");plt.title("Composición etaria de la muestra");[plt.text(x.get_x()+x.get_width()/2,v+.5,str(v),ha="center") for x,v in zip(b,an)];plt.tight_layout();plt.savefig("competencias-utmach/report_assets/edades.png",dpi=180,bbox_inches="tight");plt.close()
# Word
d=Document(); s=d.sections[0];s.top_margin=Inches(.55);s.bottom_margin=Inches(.55);s.left_margin=Inches(.65);s.right_margin=Inches(.65);d.styles["Normal"].font.name="Aptos";d.styles["Normal"].font.size=Pt(9.5)
p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run("INFORME EJECUTIVO\nDIAGNÓSTICO DE COMPETENCIAS DIGITALES DOCENTES");r.bold=True;r.font.size=Pt(22)
p=d.add_paragraph("Universidad Técnica de Machala · CEDIA");p.alignment=WD_ALIGN_PARAGRAPH.CENTER
t=d.add_table(rows=2,cols=5);t.alignment=WD_TABLE_ALIGNMENT.CENTER
for j,(v,l) in enumerate(zip(["400","136","34,0%","95%","±6,84%"],["Universo","Muestra","Cobertura","Confianza","Error"])):
 t.cell(0,j).text=v;t.cell(1,j).text=l
 for q in t.cell(0,j).paragraphs:q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.runs[0].bold=True;q.runs[0].font.size=Pt(15)
 for q in t.cell(1,j).paragraphs:q.alignment=WD_ALIGN_PARAGRAPH.CENTER
d.add_heading("1. Síntesis ejecutiva",1);d.add_paragraph("El diagnóstico institucional evidencia un perfil predominantemente Innovador en las cinco competencias. Pedagógica alcanza 75,00% en Innovador, seguida de Tecnológica (72,79%) e Investigativa (68,38%). Gestión (58,82%) y Comunicativa (58,09%) mantienen mayoría Innovador, pero concentran mayores proporciones en Explorador y constituyen las principales áreas de fortalecimiento.")
d.add_picture("competencias-utmach/report_assets/pentagono.png",width=Inches(5.7));d.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
d.add_heading("2. Alcance metodológico y representatividad",1);d.add_paragraph("El universo institucional es de 400 docentes. La base analítica final comprende 136 docentes con respuestas completas en los 45 ítems, equivalente al 34,0% del universo. Bajo corrección por población finita y p=q=0,5, el resultado agregado representa aproximadamente un 95% de confianza y ±6,84% de margen de error. Estos parámetros no deben trasladarse directamente a los subgrupos etarios.")
d.add_heading("3. Resultados institucionales por competencia",1)
tb=d.add_table(rows=1,cols=5);tb.style="Table Grid"; hdr=["Competencia","Explorador","Integrador","Innovador","Promedio /5"]
for j,x in enumerate(hdr):tb.rows[0].cells[j].text=x
for k,n in enumerate(names):
 c=tb.add_row().cells
 for j,x in enumerate([n,f"{E[k]:.2f}%",f"{I[k]:.2f}%",f"{N[k]:.2f}%",f"{A[k]:.2f}"]):c[j].text=x.replace(".",",")
d.add_picture("competencias-utmach/report_assets/distribucion.png",width=Inches(7));d.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
d.add_picture("competencias-utmach/report_assets/promedios.png",width=Inches(6.7));d.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
d.add_heading("4. Fortalezas institucionales",1);d.add_paragraph("Pedagógica: 75,00% Innovador y 4,25/5. Tecnológica: 72,79% Innovador y 4,30/5, el promedio más alto. Investigativa: 68,38% Innovador y 4,21/5. Estas dimensiones constituyen la base para estrategias de consolidación, mentoría y transferencia de buenas prácticas.")
d.add_heading("5. Brechas y prioridades",1);d.add_paragraph("Gestión constituye la principal prioridad: 31,62% en Explorador y 3,93/5. Comunicativa registra 24,26% en Explorador y 17,65% en Integrador. Investigativa mantiene 24,26% en Explorador pese a su perfil global favorable. Se recomienda focalizar formación, acompañamiento y proyectos aplicados en estas brechas.")
d.add_heading("6. Análisis por grupos etarios",1);d.add_picture("competencias-utmach/report_assets/edades.png",width=Inches(6.4));d.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER;d.add_paragraph("La muestra se distribuye en 24–35 años (25), 36–45 (47), 46–55 (42), 56–65 (19) y >66 (3). La edad se utiliza como segmentación descriptiva; no implica causalidad. El grupo >66 debe interpretarse con especial cautela por n=3.")
d.add_heading("7. Recomendaciones para toma de decisiones",1)
for x in ["Implementar rutas diferenciadas Explorador → Integrador → Innovador.","Priorizar Gestión y Comunicativa.","Consolidar fortalezas mediante mentoría y comunidades de práctica.","Vincular la formación con productos verificables y proyectos reales.","Repetir periódicamente el diagnóstico para medir evolución.","Actualizar el marco CEDIA preservando comparabilidad histórica e incorporando IA transversal."]:d.add_paragraph(x,style="List Bullet")
d.add_heading("8. Inteligencia artificial: escenario metodológico",1);d.add_paragraph("Los porcentajes de IA del visor son una simulación metodológica y no corresponden a respuestas reales de UTMACH. No forman parte de los resultados observados. La IA se plantea transversalmente sobre las cinco competencias y requerirá pilotaje y validación antes de incorporarse al diagnóstico institucional.")
d.add_heading("9. Conclusión ejecutiva",1);d.add_paragraph("UTMACH presenta una base favorable de competencias digitales docentes, con predominio Innovador en todas las dimensiones. El siguiente ciclo debe reducir brechas en Gestión y Comunicativa, sostener fortalezas Pedagógica y Tecnológica, profundizar Investigación y preparar una transición controlada hacia un marco actualizado con IA transversal.")
for sec in d.sections:
 p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run("CEDIA · Competencias Digitales Docentes · UTMACH 2026").font.size=Pt(8)
d.save("competencias-utmach/Informe_Ejecutivo_Competencias_Digitales_UTMACH_CEDIA.docx")
